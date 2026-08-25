#!/usr/bin/env python3
"""Extract manuscript structure from a .docx for paper-agent Review/Revise/Proofread/Audit modes.

Replaces the old `/mnt/skills/public/docx` delegation, which only exists inside the
Claude.ai analysis container. Runs locally on pandoc + python-docx.

Usage:
    .venv/bin/python scripts/extract_docx.py <manuscript.docx> [--sections] [--json]

Default output is the extraction report SKILL.md requires. --sections additionally
dumps the full text of every section. --json emits the whole structure as JSON.
"""
import argparse
import json
import os
import re
import sys
from collections import OrderedDict

try:
    from docx import Document
except ImportError:
    sys.exit("python-docx missing. Run: .venv/bin/pip install python-docx")

# (Author Year) / (Author, Year) / (Author et al., Year) / Author (Year)
CITE_RE = re.compile(
    r"\(([^()]*?\b(?:1[89]|20)\d{2}[a-z]?)\)"           # parenthetical
    r"|([A-Z][A-Za-zÀ-ɏ'\-]+(?:\s+(?:et\s+al\.|and\s+[A-Z][A-Za-z'\-]+))?)\s+\((\d{4}[a-z]?)\)"
)
YEAR_RE = re.compile(r"\b((?:1[89]|20)\d{2})[a-z]?\b")
FIG_REF_RE = re.compile(r"\bFig(?:ure)?s?\.?\s*(\d+)", re.I)
TAB_REF_RE = re.compile(r"\bTables?\s*(\d+)", re.I)
FIG_CAP_RE = re.compile(r"^\s*Fig(?:ure)?\.?\s*(\d+)\s*[.:—-]", re.I)
TAB_CAP_RE = re.compile(r"^\s*Table\s*(\d+)\s*[.:—-]", re.I)
EQ_NUM_RE = re.compile(r"\(\s*(\d+)\s*\)\s*$")
# A reference-list line starts with a surname and contains a year.
REF_ENTRY_RE = re.compile(r"^[A-ZÀ-ɏ][A-Za-zÀ-ɏ'\-]+,?\s")
# DOIs as they appear in reference lists: bare, doi: prefixed, or as a URL.
DOI_RE = re.compile(r"\b10\.\d{4,9}/[-._;()/:A-Za-z0-9]*[A-Za-z0-9]", re.I)

SECTION_STYLES = ("Heading 1", "Heading 2", "Heading 3", "Title")
REF_HEADINGS = ("references", "reference list", "bibliography", "works cited", "literature cited")


def is_heading(p):
    name = (p.style.name or "") if p.style is not None else ""
    if name in SECTION_STYLES or name.startswith("Heading"):
        return True
    # pandoc-produced files sometimes carry no heading style; fall back to an
    # all-bold short paragraph that looks like a numbered section head.
    txt = p.text.strip()
    if txt and len(txt) < 90 and p.runs and all(r.bold for r in p.runs if r.text.strip()):
        return bool(re.match(r"^(\d+\.?\d*\.?\d*\s+\S|Abstract|Highlights|Keywords)", txt, re.I))
    return False


def parse(path):
    doc = Document(path)
    sections = OrderedDict()
    current = "FRONT MATTER"
    sections[current] = []
    fig_caps, tab_caps, equations = [], [], []

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if is_heading(p):
            current = txt
            sections.setdefault(current, [])
            continue
        m = FIG_CAP_RE.match(txt)
        if m:
            fig_caps.append((int(m.group(1)), txt))
        m = TAB_CAP_RE.match(txt)
        if m:
            tab_caps.append((int(m.group(1)), txt))
        if EQ_NUM_RE.search(txt) and len(txt) < 400:
            equations.append(txt)
        sections[current].append(txt)

    # Reference section: split entries.
    ref_key = next((k for k in sections if k.strip().lower().rstrip(".0123456789 ") in REF_HEADINGS
                    or k.strip().lower() in REF_HEADINGS), None)
    refs = []
    if ref_key:
        raw = [ln for ln in sections[ref_key] if REF_ENTRY_RE.match(ln) and YEAR_RE.search(ln)]
        # Some exports collapse the whole list into one paragraph. Split on the
        # start of the next "Surname, Initial." entry rather than reporting 1.
        for ln in raw:
            parts = re.split(r"(?<=[.\)])\s+(?=[A-ZÀ-ɏ][A-Za-zÀ-ɏ'\-]+,\s*[A-Z])", ln)
            refs.extend(p.strip() for p in parts if p.strip() and YEAR_RE.search(p))

    body_keys = [k for k in sections if k != ref_key]
    body_text = "\n".join(" ".join(sections[k]) for k in body_keys)

    # In-text citations -> set of (surname_token, year)
    cites = []
    for m in CITE_RE.finditer(body_text):
        if m.group(1):
            inner = m.group(1)
            for chunk in re.split(r";", inner):
                y = YEAR_RE.search(chunk)
                if not y:
                    continue
                name = re.split(r",|\bet al\b|\band\b|&", chunk.strip())[0]
                # Strip the year before taking the surname: in the no-comma form
                # "(Smith 2019)" the last token is the year, not the author.
                name = YEAR_RE.sub("", name).strip(" .;")
                if name:
                    cites.append((name.split()[-1], y.group(1)))
        elif m.group(2):
            name = re.split(r"\s+et\s+al\.|\s+and\s+", m.group(2))[0].strip()
            cites.append((name.split()[-1], m.group(3)[:4]))
    cites = sorted(set(cites))

    # Match each citation to a reference entry.
    orphans = []
    used_refs = set()
    for surname, year in cites:
        hit = None
        for i, r in enumerate(refs):
            if surname.lower() in r.lower() and year in r:
                hit = i
                break
        if hit is None:
            orphans.append(f"{surname} {year}")
        else:
            used_refs.add(hit)
    unused = [refs[i][:80] for i in range(len(refs)) if i not in used_refs]

    # Figure / table cross-reference integrity.
    fig_nums = {n for n, _ in fig_caps}
    tab_nums = {n for n, _ in tab_caps}
    fig_refd = {int(n) for n in FIG_REF_RE.findall(body_text)}
    tab_refd = {int(n) for n in TAB_REF_RE.findall(body_text)}

    # Pull DOIs out of the reference entries so they can be batch-verified against
    # a live index (OpenAlex batch_resolve_references). Entries with no DOI have to
    # be verified by title match instead.
    ref_dois, no_doi = [], []
    for r in refs:
        m = DOI_RE.search(r)
        if m:
            ref_dois.append(m.group(0).rstrip(".,;"))
        else:
            no_doi.append(r[:100])

    def wc(lines):
        return sum(len(l.split()) for l in lines)

    return {
        "file": os.path.basename(path),
        "sections": OrderedDict((k, {"words": wc(v), "paragraphs": len(v), "text": v})
                                for k, v in sections.items()),
        "word_count": sum(wc(v) for k, v in sections.items() if k != ref_key),
        "figures": sorted(fig_nums),
        "tables": sorted(tab_nums),
        "fig_captions": [c for _, c in sorted(fig_caps)],
        "tab_captions": [c for _, c in sorted(tab_caps)],
        "equations": len(equations),
        "docx_tables": len(doc.tables),
        "citations": [f"{s} {y}" for s, y in cites],
        "references": refs,
        "reference_dois": ref_dois,
        "references_without_doi": no_doi,
        "ref_heading": ref_key,
        "integrity": {
            "orphan_citations": orphans,
            "unused_references": unused,
            "figs_captioned_not_cited": sorted(fig_nums - fig_refd),
            "figs_cited_not_captioned": sorted(fig_refd - fig_nums),
            "tables_captioned_not_cited": sorted(tab_nums - tab_refd),
            "tables_cited_not_captioned": sorted(tab_refd - tab_nums),
        },
    }


def report(d):
    out = [f"Manuscript loaded: {d['file']}"]
    names = list(d["sections"].keys())
    out.append(f"Sections: {', '.join(names)}")
    per = ", ".join(f"{k}: {v['words']}" for k, v in d["sections"].items() if v["words"])
    out.append(f"Word count: {d['word_count']} ({per})")
    out.append(f"Figures: {len(d['figures'])}  |  Tables: {len(d['tables'])} "
               f"(docx table objects: {d['docx_tables']})  |  Equations: {d['equations']}")
    out.append(f"Citations in text: {len(d['citations'])}  |  Reference entries: {len(d['references'])}")
    out.append(f"DOIs in reference list: {len(d['reference_dois'])} "
               f"({len(d['references_without_doi'])} entries without a DOI)")
    ig = d["integrity"]
    probs = []
    if ig["orphan_citations"]:
        probs.append(f"{len(ig['orphan_citations'])} orphan citations ({', '.join(ig['orphan_citations'][:8])}"
                     + (", ..." if len(ig["orphan_citations"]) > 8 else "") + ")")
    if ig["unused_references"]:
        probs.append(f"{len(ig['unused_references'])} unused references")
    for key, label in (("figs_cited_not_captioned", "figures cited but not captioned"),
                       ("figs_captioned_not_cited", "figures captioned but never cited"),
                       ("tables_cited_not_captioned", "tables cited but not captioned"),
                       ("tables_captioned_not_cited", "tables captioned but never cited")):
        if ig[key]:
            probs.append(f"{label}: {ig[key]}")
    out.append("Integrity: " + ("OK" if not probs else "; ".join(probs)))
    if not d["references"]:
        out.append("WARNING: no reference list parsed. Citation integrity is UNVERIFIED, "
                   "not clean. Check the References heading style before trusting the counts above.")
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--sections", action="store_true", help="dump full text of each section")
    ap.add_argument("--json", action="store_true", help="emit full structure as JSON")
    a = ap.parse_args()

    if not os.path.exists(a.docx):
        sys.exit(f"File not found: {a.docx}")
    if a.docx.lower().endswith(".doc"):
        sys.exit("Legacy .doc. Convert first:  pandoc in.doc -o in.docx")

    d = parse(a.docx)
    if a.json:
        print(json.dumps(d, indent=2, ensure_ascii=False))
        return
    print(report(d))
    if a.sections:
        for k, v in d["sections"].items():
            print(f"\n{'=' * 70}\n## {k}  ({v['words']} words)\n{'=' * 70}")
            for line in v["text"]:
                print(line + "\n")


if __name__ == "__main__":
    main()
