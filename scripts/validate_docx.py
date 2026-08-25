#!/usr/bin/env python3
"""Validate a paper-agent .docx export against the manuscript-docx-style.md spec.

Replaces the old `scripts/office/validate.py` reference, which does not exist on
this machine (it is a Claude.ai analysis-container path).

Usage:
    .venv/bin/python scripts/validate_docx.py <manuscript.docx> [--font "Times New Roman"]

Exit code 0 = all checks passed, 1 = at least one FAIL.
"""
import argparse
import os
import re
import sys
import zipfile
from collections import Counter

try:
    from docx import Document
except ImportError:
    sys.exit("python-docx missing. Run: .venv/bin/pip install python-docx")

PLACEHOLDER_RE = re.compile(r"\[(CITATION NEEDED|VALUE NEEDED|VERIFY|FROM USER)[^\]]*\]", re.I)
# Case-insensitive for I/we/my/our; "us" only in lowercase so the country
# abbreviation "US" is not flagged as a pronoun.
FIRST_PERSON_RE = re.compile(r"\b(?:[Ii]|[Ww]e|[Mm]y|[Oo]ur|[Ww]e're|[Ww]e've)\b|\bus\b")
NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

results = []


def check(name, ok, detail=""):
    results.append((name, ok, detail))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--font", default="Times New Roman", help="expected body font")
    ap.add_argument("--allow-first-person", action="store_true")
    a = ap.parse_args()

    if not os.path.exists(a.docx):
        sys.exit(f"File not found: {a.docx}")

    # 1. Opens as a valid OOXML package.
    try:
        with zipfile.ZipFile(a.docx) as z:
            bad = z.testzip()
            names = z.namelist()
        check("Valid .docx zip package", bad is None, f"corrupt member: {bad}" if bad else "")
        check("Contains word/document.xml", "word/document.xml" in names)
    except zipfile.BadZipFile:
        check("Valid .docx zip package", False, "not a zip archive")
        report()
        return

    try:
        doc = Document(a.docx)
    except Exception as e:
        check("python-docx can open the file", False, str(e))
        report()
        return
    check("python-docx can open the file", True)

    paras = [p for p in doc.paragraphs if p.text.strip()]
    body = "\n".join(p.text for p in paras)

    # 2. Non-empty.
    check("Document has body content", len(paras) > 0, f"{len(paras)} non-empty paragraphs")

    # 3. Font coercion across every run variant.
    fonts = Counter()
    for p in doc.paragraphs:
        for r in p.runs:
            if r.text.strip():
                fonts[r.font.name] += 1
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.text.strip():
                            fonts[r.font.name] += 1
    off = {f: n for f, n in fonts.items() if f not in (a.font, None)}
    check(f"Body font is {a.font} (or inherited)", not off,
          f"stray fonts: {dict(list(off.items())[:5])}" if off else
          f"{fonts.get(a.font, 0)} explicit runs, {fonts.get(None, 0)} inherited")

    # 4. Heading styles present (decimal numbered headings need real Heading styles).
    heads = [p for p in doc.paragraphs if (p.style.name or "").startswith("Heading")]
    check("Uses Heading styles", len(heads) > 0, f"{len(heads)} headings")

    # 5. Sequential figure/table caption numbering, no gaps.
    for label, rx in (("Figure", re.compile(r"^\s*Fig(?:ure)?\.?\s*(\d+)", re.I)),
                      ("Table", re.compile(r"^\s*Table\s*(\d+)", re.I))):
        nums = sorted({int(m.group(1)) for p in paras if (m := rx.match(p.text))})
        if nums:
            expected = list(range(1, len(nums) + 1))
            check(f"{label} numbering sequential from 1", nums == expected, f"found {nums}")

    # 6. Equation numbering sequential.
    eqs = sorted({int(m.group(1)) for p in paras
                  if (m := re.search(r"\(\s*(\d+)\s*\)\s*$", p.text)) and len(p.text) < 400})
    if eqs:
        check("Equation numbering sequential from 1", eqs == list(range(1, len(eqs) + 1)), f"found {eqs}")

    # 7. Unresolved placeholders surfaced, not silently shipped.
    ph = PLACEHOLDER_RE.findall(body)
    check("No unresolved placeholders", not ph,
          f"{len(ph)} found: {sorted(set(ph))[:6]}" if ph else "")

    # 8. First-person pronouns (skill rule: zero anywhere).
    if not a.allow_first_person:
        fp = FIRST_PERSON_RE.findall(body)
        check("No first-person pronouns", not fp,
              f"{len(fp)} found: {Counter(fp).most_common(5)}" if fp else "")

    # 9. Tables set to full width (Springer style, per manuscript-docx-style.md).
    if doc.tables:
        narrow = 0
        for t in doc.tables:
            el = t._tbl.find(f"{NS}tblPr")
            w = el.find(f"{NS}tblW") if el is not None else None
            if w is None or w.get(f"{NS}type") != "pct":
                narrow += 1
        check("Tables use percentage width", narrow == 0,
              f"{narrow}/{len(doc.tables)} tables lack pct width" if narrow else f"{len(doc.tables)} tables")

    # 10. Embedded images present if figures are captioned.
    media = [n for n in names if n.startswith("word/media/")]
    fig_caps = [p for p in paras if re.match(r"^\s*Fig(?:ure)?\.?\s*\d+", p.text, re.I)]
    if fig_caps:
        check("Figure images embedded", len(media) > 0,
              f"{len(fig_caps)} captions, {len(media)} embedded media files")

    report()


def report():
    fails = [r for r in results if not r[1]]
    width = max(len(n) for n, _, _ in results) + 2
    print(f"Validating export against manuscript-docx-style.md\n{'-' * (width + 30)}")
    for name, ok, detail in results:
        print(f"{'PASS' if ok else 'FAIL'}  {name.ljust(width)}{detail}")
    print(f"{'-' * (width + 30)}\n{len(results) - len(fails)}/{len(results)} checks passed.")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
